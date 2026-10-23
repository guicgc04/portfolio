"""Performance operacional turno a turno: indicadores de hoje e ontem por turno, a lista de
lotes com prazo e status, e a base para o fluxo de pendências entre turnos.

Os lotes são inventados uma a uma (criação, prazo, conclusão) e os SLAs saem da
agregação delas, como no ETL real.
"""
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from comum import DEMOS, HOJE, AGORA, INDICADORES, TURNOS, rng, salvar_json, turno_da_hora

VOLUMES = {1: 240, 2: 150, 3: 60, 4: 45, 5: 30, 6: 25, 7: 260, 8: 180, 9: 120, 10: 90, 11: 20, 12: 70}
PRAZOS = {1: 4, 2: 3, 3: 2, 4: 3, 5: 2, 6: 6, 7: 3, 8: 2, 9: 2, 10: 1, 11: 8, 12: 10}
QUALIDADE = {1: .92, 2: .90, 3: .87, 4: .93, 5: .91, 6: .95, 7: .95, 8: .955, 9: .93, 10: .965, 11: .88, 12: .94}
SLAS = [  # (item, setor, indicador, meta, lotes/dia, prazo em horas, tipos, qualidade)
    (item, setor, nome, meta, VOLUMES[item], PRAZOS[item],
     ["Manual"] if item in (4, 11) else (["Mecanizado", "Automatizado"] if item in (2, 3, 7) else ["Mecanizado", "Manual"]),
     QUALIDADE[item])
    for item, setor, nome, meta in INDICADORES
]
ORIGENS = [f"DOCA-{n:02d}" for n in range(1, 9)] + ["ESTEIRA", "PÁTIO"]
DESTINOS = [f"RUA-{c}{n:02d}" for c in "ABCD" for n in range(1, 9)] + ["DESPACHO", "QUALIDADE", "AUTOMATIZADO"]


def turno(t):
    return turno_da_hora(t.hour)


def peso_hora(h):
    return .5 if h < 6 else (1.3 if 7 <= h <= 18 else .8)


def main():
    r = rng(2207)
    dias = {"anteontem": HOJE - dt.timedelta(2), "ontem": HOJE - dt.timedelta(1), "hoje": HOJE}
    cargas, n_lote = [], [100000]
    for item, setor, nome, meta, vol, prazo_h, tipos, q in SLAS:
        if not vol:
            continue
        for rotulo, dia in dias.items():
            ini_h = 22 if rotulo == "anteontem" else 0
            for h in range(ini_h, 24):
                inicio_hora = dt.datetime.combine(dia, dt.time(h))
                if inicio_hora > AGORA:
                    break
                n = vol / 24 * peso_hora(h) * r.uniform(.6, 1.4)
                for _ in range(int(n) + (r.random() < n - int(n))):
                    criacao = inicio_hora + dt.timedelta(minutes=r.randint(0, 59))
                    if criacao > AGORA:
                        continue
                    vence = criacao + dt.timedelta(hours=prazo_h)
                    no_prazo = r.random() < q
                    dur = dt.timedelta(minutes=r.randint(5, int(prazo_h * 60 * .95))) if no_prazo else \
                        dt.timedelta(minutes=int(prazo_h * 60 * r.uniform(1.05, 1.8)))
                    concl = criacao + dur
                    concluido = concl <= AGORA
                    if concluido:
                        status = "dentro" if concl <= vence else "vencido"
                    else:
                        falta = (vence - AGORA).total_seconds() / 3600
                        status = "vencido" if falta < 0 else ("prestes" if falta < .5 else "dentro")
                    n_lote[0] += 1
                    c = {"item": item, "indicador": nome, "lote": f"LT-{n_lote[0]}",
                         "pedido": "" if item == 1 else f"PED-{r.randint(1000000, 9999999)}",
                         "origem": r.choice(ORIGENS), "destino": "" if item == 1 else r.choice(DESTINOS),
                         "tipo": r.choice(tipos), "criacao": criacao.isoformat(timespec="seconds"),
                         "conclusao": concl.isoformat(timespec="seconds") if concluido else None,
                         "vence_em": vence.isoformat(timespec="seconds"),
                         "turno": turno(concl if concluido else AGORA), "turno_origem": turno(criacao),
                         "turno_fim": turno(concl) if concluido else None, "concluido": concluido,
                         "status": status, "dia": rotulo}
                    if rotulo == "anteontem":
                        c["so_fluxo"] = True
                    cargas.append(c)

    def slas_do_dia(dia):
        saida = []
        for item, setor, nome, meta, *_ in SLAS:
            feitas = [c for c in cargas if c["item"] == item and c["concluido"]
                      and dt.datetime.fromisoformat(c["conclusao"]).date() == dia]
            turnos = {t: {"dentro": sum(1 for c in feitas if c["turno_fim"] == t and c["status"] == "dentro"),
                          "fora": sum(1 for c in feitas if c["turno_fim"] == t and c["status"] != "dentro")}
                      for t in TURNOS}
            dentro = sum(v["dentro"] for v in turnos.values())
            fora = sum(v["fora"] for v in turnos.values())
            total = dentro + fora
            saida.append({"item": item, "setor": setor, "indicador": nome, "dentro": dentro, "fora": fora,
                          "total": total, "meta": meta, "sla_pct": round(dentro / total, 4) if total else None,
                          "turnos": turnos})
        return saida

    pend = []
    for item, setor, nome, *_ in SLAS:
        abertas = [c for c in cargas if c["item"] == item and not c["concluido"]]
        pend.append({"item": item, "setor": setor, "indicador": nome, "total_pendente": len(abertas),
                     "turnos": {t: sum(1 for c in abertas if c["turno_origem"] == t) for t in TURNOS},
                     "sem_dado": False})

    salvar_json("performance-turnos", {
        "atualizado_em": AGORA.isoformat(),
        "hoje": {"data": HOJE.isoformat(), "slas": slas_do_dia(HOJE)},
        "ontem": {"data": dias["ontem"].isoformat(), "slas": slas_do_dia(dias["ontem"])},
        "cargas": cargas,
        "pendencias_turno": pend,
    }, "dados_dashboard.json")
    print(f"  {len(cargas)} lotes fictícios")
    exportar_planilha([c for c in cargas if not c.get("so_fluxo")])


def exportar_planilha(cargas):
    """Planilha do botão "Baixar" da aba Cargas (também fictítransp)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Lotes"
    colunas = ["dia", "item", "indicador", "lote", "pedido", "origem", "destino", "tipo",
               "criacao", "conclusao", "vence_em", "turno", "status"]
    ws.append(["Dia", "Item", "Indicador", "Lote", "Pedido", "Origem", "Destino", "Tipo", "Criação", "Conclusão", "Vence em", "Turno", "Status"])
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="2B3140")
    for c in cargas:
        ws.append([c.get(k) for k in colunas])
    ws.freeze_panes = "A2"
    for col, larg in zip("ABCDEFGHIJKLM", [8, 6, 34, 12, 14, 10, 12, 14, 20, 20, 20, 9, 9]):
        ws.column_dimensions[col].width = larg
    wb.save(os.path.join(DEMOS, "performance-turnos", "detalhe_cargas.xlsx"))


if __name__ == "__main__":
    main()
